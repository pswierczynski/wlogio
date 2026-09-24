"""
export.py - Eksport danych panelu pracy do CSV / XLSX / PDF.

Eksportuje wyłącznie to, co użytkownik widzi w swoim panelu pracy — te same
obliczenia (godziny, wynagrodzenie, nadgodziny, stawka za dni poza roboczymi)
co dashboard, poprzez współdzieloną funkcję build_months_data().
"""
import csv
import io
import os
from datetime import date, timedelta

from flask import Blueprint, render_template, request, Response, flash, redirect, url_for, current_app
from flask_login import login_required, current_user

from wlogio_app.routes.dashboard import build_months_data
from wlogio_app.calculator import (
    format_currency,
    _get_config_value,
    _parse_work_days,
    DEFAULT_OVERTIME_RATE,
    DEFAULT_OFFDAY_RATE,
    DEFAULT_HOURS_PER_DAY,
)

export_bp = Blueprint('export', __name__)

TYPE_LABELS = {
    'work':          'Praca',
    'vacation':      'Urlop',
    'on_demand':     'Urlop na żądanie',
    'unpaid':        'Bezpłatny',
    'holiday':       'Święto',
    'sick_leave':    'Zwolnienie lekarskie',
    'care_leave':    'Urlop opiekuńczy',
    'force_majeure': 'Siła wyższa',
    'child_care':    'Opieka nad dzieckiem',
}

RANGE_LABELS = {
    'custom':         'Zakres własny',
    'current_year':   'Bieżący rok (do dziś)',
    'previous_year':  'Poprzedni rok',
    'previous_month': 'Poprzedni miesiąc',
}


def entry_type_label(entry):
    label = TYPE_LABELS.get(entry.entry_type, entry.entry_type)
    if entry.entry_type in ('vacation', 'on_demand') and entry.vacation_day_number:
        label += f' ({entry.vacation_day_number})'
    if entry.is_remote:
        label += ' — zdalna'
    return label


def entry_break_minutes(entry):
    """Suma minut wszystkich przerw wpisu (parsowanie identyczne jak w panelu)."""
    if not entry.breaks:
        return 0
    total = 0
    for seg in entry.breaks.split(';'):
        seg = seg.strip()
        if seg and '-' in seg:
            parts = seg.split('-')
            if len(parts) == 2 and parts[0] and parts[1]:
                try:
                    sh, sm = map(int, parts[0].split(':'))
                    eh, em = map(int, parts[1].split(':'))
                    total += (eh * 60 + em) - (sh * 60 + sm)
                except ValueError:
                    continue
    return total


def entry_salary(entry, month):
    """
    Wynagrodzenie za pojedynczy dzień — dokładnie ta sama logika co w
    wierszu tabeli panelu pracy (dashboard/index.html): stawka za dni poza
    roboczymi dla weekendów/dni wolnych, stawka + nadgodziny dla dni roboczych.
    Zwraca None gdy nie da się policzyć (brak stawki, dzień bezpłatny, albo
    wpis w trakcie trwania — brak godziny zakończenia).
    """
    hourly_rate = month['hourly_rate']
    if not hourly_rate or hourly_rate <= 0 or entry.entry_type == 'unpaid':
        return None
    if entry.entry_type == 'work' and entry.time_start and not entry.time_end:
        return None

    config = month['config']
    hpd = float(_get_config_value(config, 'hours_per_day', DEFAULT_HOURS_PER_DAY))

    if entry.entry_type == 'work':
        ot_rate = _get_config_value(config, 'overtime_rate', DEFAULT_OVERTIME_RATE) / 100
        od_rate = _get_config_value(config, 'offday_rate', DEFAULT_OFFDAY_RATE) / 100
        work_days = _get_config_value(config, 'work_days', None)
        work_days_list = _parse_work_days(work_days) if work_days else [0, 1, 2, 3, 4]
        is_offday = entry.date.weekday() not in work_days_list

        billed = float(entry.hours_billed)
        if is_offday:
            return billed * hourly_rate * od_rate
        normal_h = min(billed, hpd)
        ot_h = max(0.0, billed - hpd)
        return normal_h * hourly_rate + ot_h * hourly_rate * ot_rate

    return float(entry.hours_billed) * hourly_rate


def resolve_date_range(range_key, date_from_str, date_to_str):
    today = date.today()
    if range_key == 'current_year':
        return date(today.year, 1, 1), today
    if range_key == 'previous_year':
        y = today.year - 1
        return date(y, 1, 1), date(y, 12, 31)
    if range_key == 'previous_month':
        first_this = date(today.year, today.month, 1)
        last_prev = first_this - timedelta(days=1)
        first_prev = date(last_prev.year, last_prev.month, 1)
        return first_prev, last_prev
    # custom
    d_from = date.fromisoformat(date_from_str) if date_from_str else None
    d_to = date.fromisoformat(date_to_str) if date_to_str else None
    return d_from, d_to


def hours_str(value):
    return f"{float(value):.2f}".replace('.', ',')


# ---------------------------------------------------------------- CSV ----

def generate_csv(months_data):
    output = io.StringIO()
    output.write('\ufeff')  # BOM — poprawne polskie znaki przy otwieraniu w Excelu
    writer = csv.writer(output, delimiter=';')

    for month in months_data:
        writer.writerow([f"{month['month_name']} {month['year']}"])
        writer.writerow(['Data', 'Typ', 'Przyjście', 'Wyjście', 'Przerwa (min)', 'Godziny', 'Wynagrodzenie'])

        for entry in month['entries']:
            salary = entry_salary(entry, month)
            bmin = entry_break_minutes(entry)
            writer.writerow([
                entry.date.strftime('%Y-%m-%d'),
                entry_type_label(entry),
                entry.time_start.strftime('%H:%M') if entry.time_start else '',
                entry.time_end.strftime('%H:%M') if entry.time_end else '',
                bmin if bmin else '',
                hours_str(entry.hours_billed) if entry.entry_type != 'unpaid' else '0',
                format_currency(salary) if salary is not None else '',
            ])

        s = month['summary']
        writer.writerow([])
        writer.writerow(['Podsumowanie miesiąca'])
        writer.writerow(['Przepracowane (h)', hours_str(s['total_hours'])])
        writer.writerow(['Wymagane (h)', f"{s['expected_hours']:.0f}"])
        writer.writerow(['Nadgodziny (h)', hours_str(s['overtime_hours'])])
        writer.writerow(['Wynagrodzenie', format_currency(s['actual_salary'])])
        if s.get('bonus'):
            writer.writerow(['Premia', format_currency(s['bonus'])])
        writer.writerow(['Razem', format_currency(s['total_with_bonus'])])
        writer.writerow(['Stawka', f"{month['hourly_rate']} PLN/h"])
        writer.writerow([])
        writer.writerow([])

    return output.getvalue().encode('utf-8')


# --------------------------------------------------------------- XLSX ----

def generate_xlsx(months_data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'Eksport wlogio'

    bold = Font(bold=True)
    month_font = Font(bold=True, size=13)
    header_fill = PatternFill(start_color='EEEEEE', end_color='EEEEEE', fill_type='solid')
    headers = ['Data', 'Typ', 'Przyjście', 'Wyjście', 'Przerwa (min)', 'Godziny', 'Wynagrodzenie']

    row = 1
    for month in months_data:
        ws.cell(row=row, column=1, value=f"{month['month_name']} {month['year']}").font = month_font
        row += 1

        for col, h in enumerate(headers, start=1):
            c = ws.cell(row=row, column=col, value=h)
            c.font = bold
            c.fill = header_fill
        row += 1

        for entry in month['entries']:
            salary = entry_salary(entry, month)
            bmin = entry_break_minutes(entry)
            ws.cell(row=row, column=1, value=entry.date.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=2, value=entry_type_label(entry))
            ws.cell(row=row, column=3, value=entry.time_start.strftime('%H:%M') if entry.time_start else '')
            ws.cell(row=row, column=4, value=entry.time_end.strftime('%H:%M') if entry.time_end else '')
            ws.cell(row=row, column=5, value=bmin if bmin else None)
            hb_cell = ws.cell(
                row=row, column=6,
                value=round(float(entry.hours_billed), 2) if entry.entry_type != 'unpaid' else 0
            )
            hb_cell.alignment = Alignment(horizontal='right')
            sal_cell = ws.cell(row=row, column=7, value=round(salary, 2) if salary is not None else None)
            sal_cell.alignment = Alignment(horizontal='right')
            if salary is not None:
                sal_cell.number_format = '#,##0.00 "zł"'
            row += 1

        s = month['summary']
        row += 1
        ws.cell(row=row, column=1, value='Podsumowanie miesiąca').font = bold
        row += 1
        ws.cell(row=row, column=1, value='Przepracowane (h)')
        ws.cell(row=row, column=2, value=round(s['total_hours'], 2))
        row += 1
        ws.cell(row=row, column=1, value='Wymagane (h)')
        ws.cell(row=row, column=2, value=round(s['expected_hours'], 2))
        row += 1
        ws.cell(row=row, column=1, value='Nadgodziny (h)')
        ws.cell(row=row, column=2, value=round(s['overtime_hours'], 2))
        row += 1
        ws.cell(row=row, column=1, value='Wynagrodzenie')
        ws.cell(row=row, column=2, value=round(s['actual_salary'], 2)).number_format = '#,##0.00 "zł"'
        row += 1
        if s.get('bonus'):
            ws.cell(row=row, column=1, value='Premia')
            ws.cell(row=row, column=2, value=round(s['bonus'], 2)).number_format = '#,##0.00 "zł"'
            row += 1
        ws.cell(row=row, column=1, value='Razem').font = bold
        razem_cell = ws.cell(row=row, column=2, value=round(s['total_with_bonus'], 2))
        razem_cell.font = bold
        razem_cell.number_format = '#,##0.00 "zł"'
        row += 1
        ws.cell(row=row, column=1, value='Stawka')
        ws.cell(row=row, column=2, value=f"{month['hourly_rate']} PLN/h")
        row += 3

    widths = [13, 22, 11, 11, 13, 11, 16]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- PDF ----

def generate_pdf(months_data, user):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    fonts_dir = os.path.join(current_app.static_folder, 'fonts')
    pdfmetrics.registerFont(TTFont('DejaVuSans', os.path.join(fonts_dir, 'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('DejaVuSans-Bold', os.path.join(fonts_dir, 'DejaVuSans-Bold.ttf')))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        topMargin=16 * mm, bottomMargin=16 * mm, leftMargin=14 * mm, rightMargin=14 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitlePL', parent=styles['Title'], fontName='DejaVuSans-Bold', fontSize=16)
    heading_style = ParagraphStyle('HeadingPL', parent=styles['Heading2'], fontName='DejaVuSans-Bold', fontSize=12)
    normal_style = ParagraphStyle('NormalPL', parent=styles['Normal'], fontName='DejaVuSans', fontSize=9, leading=13)

    elements = [
        Paragraph(f"Eksport czasu pracy — {user.name}", title_style),
        Spacer(1, 12),
    ]

    headers = ['Data', 'Typ', 'Przyjście', 'Wyjście', 'Przerwa', 'Godziny', 'Wynagrodzenie']

    for month in months_data:
        elements.append(Paragraph(f"{month['month_name']} {month['year']}", heading_style))
        elements.append(Spacer(1, 4))

        data = [headers]
        for entry in month['entries']:
            salary = entry_salary(entry, month)
            bmin = entry_break_minutes(entry)
            data.append([
                entry.date.strftime('%d.%m.%Y'),
                entry_type_label(entry),
                entry.time_start.strftime('%H:%M') if entry.time_start else '—',
                entry.time_end.strftime('%H:%M') if entry.time_end else '—',
                f'{bmin}min' if bmin else '—',
                (hours_str(entry.hours_billed) + 'h') if entry.entry_type != 'unpaid' else '0h',
                format_currency(salary) if salary is not None else '—',
            ])

        table = Table(
            data, repeatRows=1,
            colWidths=[20 * mm, 34 * mm, 16 * mm, 16 * mm, 16 * mm, 18 * mm, 28 * mm],
        )
        table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, 0), 'DejaVuSans-Bold'),
            ('FONTNAME', (0, 1), (-1, -1), 'DejaVuSans'),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#eeeeee')),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#cccccc')),
            ('ALIGN', (5, 0), (6, -1), 'RIGHT'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        elements.append(table)

        s = month['summary']
        summary_text = (
            f"<b>Przepracowane:</b> {hours_str(s['total_hours'])}h &nbsp;&nbsp; "
            f"<b>Wymagane:</b> {s['expected_hours']:.0f}h &nbsp;&nbsp; "
            f"<b>Nadgodziny:</b> {hours_str(s['overtime_hours'])}h &nbsp;&nbsp; "
            f"<b>Wynagrodzenie:</b> {format_currency(s['actual_salary'])} &nbsp;&nbsp; "
            f"<b>Razem:</b> {format_currency(s['total_with_bonus'])} &nbsp;&nbsp; "
            f"<b>Stawka:</b> {month['hourly_rate']} PLN/h"
        )
        elements.append(Spacer(1, 6))
        elements.append(Paragraph(summary_text, normal_style))
        elements.append(Spacer(1, 18))

    if len(months_data) == 0 or all(len(m['entries']) == 0 for m in months_data):
        elements.append(Paragraph("Brak wpisów w wybranym zakresie dat.", normal_style))

    doc.build(elements)
    return buf.getvalue()


# --------------------------------------------------------------- ROUTES ----

@export_bp.route('/', methods=['GET'])
@login_required
def index():
    return render_template('export/index.html')


@export_bp.route('/', methods=['POST'])
@login_required
def generate():
    fmt = request.form.get('format', 'csv')
    range_key = request.form.get('range', 'current_year')
    date_from_str = request.form.get('date_from')
    date_to_str = request.form.get('date_to')

    if fmt not in ('csv', 'xlsx', 'pdf'):
        flash('Nieznany format eksportu.', 'error')
        return redirect(url_for('export.index'))

    if range_key == 'custom' and (not date_from_str or not date_to_str):
        flash('Podaj obie daty zakresu (od — do).', 'error')
        return redirect(url_for('export.index'))

    date_from, date_to = resolve_date_range(range_key, date_from_str, date_to_str)

    if date_from and date_to and date_from > date_to:
        flash('Data "od" musi być wcześniejsza niż data "do".', 'error')
        return redirect(url_for('export.index'))

    months_data, _ = build_months_data(
        current_user.id, date_from=date_from, date_to=date_to, include_empty_current=False
    )

    filename_base = f"wlogio_eksport_{date.today().isoformat()}"

    if fmt == 'csv':
        content = generate_csv(months_data)
        return Response(
            content, mimetype='text/csv; charset=utf-8',
            headers={'Content-Disposition': f'attachment; filename={filename_base}.csv'},
        )
    if fmt == 'xlsx':
        content = generate_xlsx(months_data)
        return Response(
            content,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            headers={'Content-Disposition': f'attachment; filename={filename_base}.xlsx'},
        )
    content = generate_pdf(months_data, current_user)
    return Response(
        content, mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename={filename_base}.pdf'},
    )
